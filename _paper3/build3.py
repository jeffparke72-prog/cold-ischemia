#!/usr/bin/env python3
"""Build DOCX, PDF and EPUB for 'What We Promised' from ms_split/*.md"""
import re, os, sys, html, zipfile, subprocess, glob, uuid
HERE=os.path.dirname(os.path.abspath(__file__)); SPL=HERE+'/ms_split'; FIG=HERE+'/fig'
OUT=sys.argv[1] if len(sys.argv)>1 else '/tmp/out/paper3'; os.makedirs(OUT,exist_ok=True)
TITLE='What We Promised'
SUB='The Kidney Patient, the Afternoon, and America’s Unfinished Pledge'
AUTHOR='Jeff Parke'
ALT={
'promise':'Bar chart comparing about 10,000 people expected to be covered by the 1972 kidney entitlement with 525,481 Medicare beneficiaries with kidney failure in 2012.',
'time':'Bar chart comparing a 2,080-hour full-time work year with a year of in-center hemodialysis: 624 hours in the chair plus 468 hours recovering, 1,092 hours in all.',
'song':'Bar chart of outcomes that patients and caregivers rated higher than clinicians: ability to travel plus 0.9, time free from dialysis plus 0.5, adequacy of dialysis plus 0.3, feeling washed out plus 0.2.',
'preg':'Two bar charts of live-birth rates in pregnancies on dialysis: 86.4 percent in a Toronto intensive-dialysis cohort versus 61.4 percent in a US registry; 48 percent at 20 hours a week or fewer versus 85 percent above 36 hours.',
'fill':'Bar chart of adult nephrology fellowship fill rates: 94.1 percent in 2010, 65.8 percent in 2024, about 73 percent in 2025.',
'chain':'Schematic of a paired-donation chain: a non-directed donor gives first, each pair donor gives to the next recipient, and the final kidney goes to a recipient on the waiting list.',
'crowd':'Bar chart of the share of requested amounts raised by Canadian transplant crowdfunding campaigns: kidney 11.5 percent, liver nearly half.',
'hazard':'Line chart of the relative hazard of graft failure and death rising with hours in cold storage from 6 to 36 hours.',
'graft':'Bar chart of median kidney graft survival in years: deceased donor 8.2 for 1995 to 1999 and 11.7 in the recent era; living donor 12.1 and 19.2.',
'nonuse':'Line chart of the share of recovered kidneys not transplanted: 18.2 percent in 2013, 26.6 in 2022, 27.9 in 2023 and 29.3 in 2024.'}
def smart(t):
    t=re.sub(r'(^|[\s(\[—–])"',lambda m:m.group(1)+'“',t); t=t.replace('"','”')
    t=re.sub(r"(^|[\s(\[—])'",lambda m:m.group(1)+'‘',t); return t.replace("'",'’')
def parse():
    B=[]
    for f in sorted(glob.glob(SPL+'/*.md')):
        L=open(f,encoding='utf-8').read().split('\n'); i=0
        while i<len(L):
            s=L[i].rstrip()
            if not s.strip(): i+=1; continue
            if s.startswith('@@PART'):
                m=re.match(r'@@PART\s+(.*?)\s*\|\s*(.*)',s); B.append(('part',m.group(1),m.group(2))); i+=1; continue
            if s.startswith('@@'):
                tag=s[2:].strip(); i+=1; body=[]
                while i<len(L) and L[i].strip() and not L[i].startswith('@@'): body.append(L[i].strip()); i+=1
                B.append(('special',tag,body)); continue
            if s.startswith('## '): B.append(('h3',s[3:].strip())); i+=1; continue
            if s.startswith('# '): B.append(('h1',s[2:].strip())); i+=1; continue
            if s.startswith('EPIGRAPH:'):
                q=s[9:].strip(); by=''
                if i+1<len(L) and L[i+1].startswith('BY:'): by=L[i+1][3:].strip(); i+=1
                B.append(('epi',q,by)); i+=1; continue
            m=re.match(r'\[\[FIG:([a-z]+)\|(.*)\]\]',s)
            if m: B.append(('fig',m.group(1),m.group(2))); i+=1; continue
            if s.startswith('|'):
                rows=[]
                while i<len(L) and L[i].startswith('|'):
                    cells=[c.strip() for c in L[i].strip().strip('|').split('|')]
                    if not all(re.fullmatch(r'-+',c) for c in cells): rows.append(cells)
                    i+=1
                B.append(('table',rows)); continue
            B.append(('p',s.strip())); i+=1
    return B
BL=parse()
def kind(txt): return 'chapter' if txt.startswith('Chapter ') else 'front'
def split_ch(txt):
    m=re.match(r'Chapter (\d+):\s*(.*)',txt); return (m.group(1),m.group(2)) if m else (None,txt)
# headings registry: (id, level, text, target_id)
HEADS=[]; n=0; pend=None
for b in BL:
    if b[0]=='part': pend=(b[1],b[2])
    if b[0]=='h1':
        n+=1
        if pend: HEADS.append((f'p{n}',1,f'Part {pend[0]}: {pend[1]}',f's{n}')); pend=None
        HEADS.append((f's{n}',2 if kind(b[1])=='chapter' else 1,b[1],f's{n}'))
def inline_html(t):
    t=html.escape(smart(t),quote=False); t=re.sub(r'\*\*(.+?)\*\*',r'<strong>\1</strong>',t)
    return re.sub(r'(?<!\*)\*(?!\s)(.+?)(?<!\s)\*(?!\*)',r'<em>\1</em>',t)
def runs(t):
    t=smart(t); out=[]
    for p in re.split(r'(\*\*.+?\*\*|(?<!\*)\*(?!\s).+?(?<!\s)\*(?!\*))',t):
        if not p: continue
        if p.startswith('**') and p.endswith('**'): out.append((p[2:-2],'b'))
        elif p.startswith('*') and p.endswith('*') and len(p)>2: out.append((p[1:-1],'i'))
        else: out.append((p,''))
    return out
COPY=[f'<strong>{TITLE}</strong>',SUB,'Copyright © 2026 Jeff Parke. All rights reserved.','Published by the Cold Ischemia Foundation, Ellenton, Florida. coldischemia.foundation','First edition, October 2026.',
'No part of this book may be reproduced, stored in a retrieval system, or transmitted in any form or by any means without the prior written permission of the publisher, except for brief quotations used in reviews or scholarly work with full attribution.',
'<strong>Educational purpose.</strong> This book describes systems and people. It is not medical, legal, or financial advice and does not replace the counsel of your own clinicians. In a medical emergency call 911. If you or someone you love is in emotional crisis, call or text 988.',
'<strong>On people and scenes.</strong> Named people appear because their stories are in the public record. Other individuals, identified by a first name or by no name, are composites built from documented patterns and are marked as such in the text.',
'<strong>On accuracy.</strong> Statements about organizations, investigations, and legal matters describe the sources cited. Sources marked with a dagger (†) have bibliographic details that should be completed from the original before any future edition. Corrections are welcome through the Foundation’s contact page.',
'<strong>Quotations.</strong> The epigraphs are drawn from works in the public domain.']
CSS='''
body{font-family:"Crimson Pro",Georgia,serif;color:#1b2430;line-height:1.58;font-size:11.2pt;text-align:left;hyphens:auto}
p{margin:0 0 .7em;text-indent:1.3em;orphans:3;widows:3} p.first{text-indent:0}
h1,h2,h3{font-family:"Playfair Display",Georgia,serif;color:#1c2b3a;hyphens:none;text-align:left}
.opener{padding-top:1.15in} .opener.nopad{padding-top:0}
.partlabel{font-family:"Bebas Neue",sans-serif;letter-spacing:5px;font-size:13pt;color:#c96a14;margin:0 0 .5in;text-align:center}
.chlabel{font-family:"Bebas Neue",sans-serif;letter-spacing:6px;font-size:12pt;color:#0e8fb0;margin:0 0 .12in}
h2.chap{font-size:24pt;line-height:1.12;margin:0 0 .3in}
h1.front{font-size:22pt;line-height:1.15;margin:0 0 .3in}
h3{font-size:13pt;margin:1.5em 0 .5em;line-height:1.25;page-break-after:avoid}
.epi{margin:.1in .4in .45in;font-style:italic;color:#3a4756;font-size:10.6pt;line-height:1.5} .epi .by{display:block;text-align:right;font-style:normal;font-size:9pt;letter-spacing:.4px;margin-top:.4em;color:#6b7684}
figure{margin:1.2em 0;text-align:center;page-break-inside:avoid} figure img{width:100%;height:auto} figcaption{font-size:9pt;color:#4a5563;text-align:left;margin-top:.5em;line-height:1.4;font-style:italic}
.note{border-left:4px solid #c96a14;background:#f6efe3;padding:.7em 1em;margin:1.1em 0;page-break-inside:avoid} .note p{text-indent:0;margin:0 0 .4em;font-size:10.4pt}
table{border-collapse:collapse;width:100%;margin:1.1em 0;font-size:9pt;text-align:left} th{background:#1c2b3a;color:#fff;padding:5px 7px;text-align:left} td{border-bottom:1px solid #d3dae2;padding:6px 7px;vertical-align:top} tr{page-break-inside:avoid}
.toc{font-size:9pt} .toc div{display:flex;align-items:baseline;margin:1.5px 0;line-height:1.28} .toc a{color:#1c2b3a;text-decoration:none} .toc .l1{margin-top:6px;font-family:"Bebas Neue",sans-serif;letter-spacing:2.5px;font-size:11pt;color:#c96a14} .toc .l2{padding-left:1.2em} .toc .l3{font-weight:700} .toc .dots{flex:1;border-bottom:1px dotted #9aa5b1;margin:0 6px;transform:translateY(-3px)}
.title{text-align:center;padding-top:2.4in} .title h1{font-size:38pt;letter-spacing:3px;text-align:center;margin:0;line-height:1.05} .title .sub{font-style:italic;font-size:13.5pt;margin:.45in .45in;color:#3a4756;line-height:1.45} .title .au{font-family:"Bebas Neue",sans-serif;letter-spacing:8px;font-size:22pt;margin-top:1in;color:#1c2b3a} .title .fd{font-size:8.5pt;letter-spacing:3px;color:#6b7684;text-transform:uppercase;margin-top:.1in}
.copy{padding-top:3.1in;font-size:9pt;line-height:1.5;text-align:left} .copy p{text-indent:0;margin-bottom:.8em}
.ded{padding-top:1.7in;text-align:center;font-style:italic;font-size:13pt;padding-left:.7in;padding-right:.7in;line-height:1.5} .kel{margin-top:2in;text-align:center;font-style:italic;font-size:11pt;padding:0 .6in;color:#3a4756} .kel .by{display:block;font-style:normal;font-size:9pt;margin-top:.5em;color:#6b7684}
.refs p{text-indent:-1.3em;margin-left:1.3em;font-size:9.4pt;text-align:left;word-break:break-word;margin-bottom:.55em}
'''
PDFONLY='.pb{page-break-before:always} .keeptogether{page-break-inside:avoid}'
def render(toc_pages=None,epub=False):
    out=[]; first=False; pend=None; refs=False; hn=0; sawtoc=False
    for b in BL:
        t=b[0]
        if t=='part': pend=(b[1],b[2]); continue
        if t=='special':
            tag,body=b[1],b[2]
            if tag=='TITLEPAGE': out.append(f'<div class="title"><h1>{TITLE.upper()}</h1><div class="sub">{html.escape(SUB)}</div><div class="au">{AUTHOR.upper()}</div><div class="fd">Cold Ischemia Foundation</div></div>')
            elif tag=='COPYRIGHT': out.append('<div class="copy pb">'+''.join(f'<p>{c}</p>' for c in COPY)+'</div>')
            elif tag=='DEDICATION': ded=' '.join(body); out.append(f'<div class="pb"><div class="ded">{inline_html(ded)}</div>')
            elif tag=='EPIGRAPH': out.append(f'<div class="kel">{inline_html(body[0])}<span class="by">{inline_html(body[1] if len(body)>1 else "")}</span></div></div>')
            elif tag=='TOC':
                rows=[]
                for hid,lvl,txt,tgt in HEADS:
                    label=html.escape(smart(txt)); pg=(toc_pages or {}).get(hid,'')
                    cls='l1' if lvl==1 and txt.startswith('Part ') else ('l2' if lvl==2 else 'l3')
                    if epub: rows.append(f'<div class="{cls}"><a href="{tgt}.xhtml">{label}</a></div>')
                    else: rows.append(f'<div class="{cls}"><a href="#{tgt}">{label}</a><span class="dots"></span><span>{pg}</span></div>')
                out.append('<div class="pb"><div class="opener nopad"><h1 class="front" id="toc">Contents</h1></div><div class="toc">'+''.join(rows)+'</div></div>')
            continue
        if t=='h1':
            hn+=1; k=kind(b[1]); refs=b[1].startswith('Appendix B')
            num,title=split_ch(b[1]); lab=''
            if pend: lab=f'<div class="partlabel">PART {html.escape(pend[0].upper())}<br/>{html.escape(smart(pend[1]).upper())}</div>'; pend=None
            pb='' if epub else ' pb'
            if k=='chapter': head=f'<div class="chlabel">CHAPTER {num}</div><h2 class="chap" id="s{hn}">{html.escape(smart(title))}</h2>'
            else: head=f'<h1 class="front" id="s{hn}">{html.escape(smart(b[1]))}</h1>'
            out.append(f'<div class="opener{pb}">{lab}{head}</div>'); first=True; continue
        if t=='h3': out.append(f'<h3>{html.escape(smart(b[1]))}</h3>'); first=True; continue
        if t=='epi': out.append(f'<div class="epi">{inline_html(b[1])}<span class="by">{inline_html(b[2])}</span></div>'); continue
        if t=='fig': out.append(f'<figure><img src="fig/{b[1]}.png" alt="{html.escape(ALT[b[1]])}"/><figcaption>{inline_html(b[2])}</figcaption></figure>'); continue
        if t=='note': out.append('<div class="note">'+''.join(f'<p>{inline_html(x)}</p>' for x in b[2])+'</div>'); continue
        if t=='table':
            r=b[1]; out.append('<table><thead><tr>'+''.join(f'<th>{inline_html(c)}</th>' for c in r[0])+'</tr></thead><tbody>'+''.join('<tr>'+''.join(f'<td>{inline_html(c)}</td>' for c in x)+'</tr>' for x in r[1:])+'</tbody></table>'); continue
        if t=='p':
            if refs: out.append(f'<div class="refs"><p>{inline_html(b[1])}</p></div>')
            else: out.append(f'<p class="{"first" if first else ""}">{inline_html(b[1])}</p>')
            first=False
    return '\n'.join(out)

def build_pdf():
    import pymupdf as fitz
    fonts=''.join(f'@font-face{{font-family:"{n}";src:url("file:///home/user/cold-ischemia/fonts/{f}");font-weight:{w};font-style:{s}}}' for n,f,w,s in [('Crimson Pro','crimsonpro-400-normal.woff2',400,'normal'),('Crimson Pro','crimsonpro-400-italic.woff2',400,'italic'),('Crimson Pro','crimsonpro-600-normal.woff2',700,'normal'),('Playfair Display','playfairdisplay-900-normal.woff2',700,'normal'),('Bebas Neue','bebasneue-400-normal.woff2',400,'normal')])
    js="const pw=require('/opt/node22/lib/node_modules/playwright');(async()=>{const b=await pw.chromium.launch();const p=await b.newPage();await p.goto('file://%s');await p.evaluate(()=>document.fonts.ready);await p.waitForTimeout(800);await p.pdf({path:'%s',width:'6in',height:'9in',printBackground:true,preferCSSPageSize:true});await b.close()})()"
    toc=None
    for _ in (1,2):
        body=render(toc)
        doc=f'<!doctype html><html><head><meta charset="utf-8"><base href="file://{HERE}/"><style>@page{{size:6in 9in;margin:0.8in 0.72in 0.85in 0.72in}}{fonts}{CSS}{PDFONLY}</style></head><body>{body}</body></html>'
        hp=f'{OUT}/_body.html'; open(hp,'w',encoding='utf-8').write(doc); bp=f'{OUT}/_body.pdf'
        open('/tmp/_pdfjs2.js','w').write(js%(hp,bp)); subprocess.run(['node','/tmp/_pdfjs2.js'],check=True)
        d=fitz.open(bp); toc={}; start=0
        for pi in range(len(d)):
            if d[pi].get_text().strip().startswith('Contents'): start=pi+1; break
        cur=start
        # find the page of every heading, in order, after the TOC pages
        while cur<len(d) and 'Contents' not in d[cur].get_text()[:40] and False: cur+=1
        for hid,lvl,txt,tgt in HEADS:
            num,title=split_ch(txt); needle=smart(title if num else txt)
            if txt.startswith('Part '): needle=None
            found=None
            for pi in range(cur,len(d)):
                if needle and any(r.y0<330 for r in d[pi].search_for(needle[:50])): found=pi; break
                if not needle: found=None; break
            if txt.startswith('Part '):
                toc[hid]=None; continue
            if found is None: print('NOT FOUND',txt); found=cur
            toc[hid]=found+2; cur=found
        # part rows take the page of the following chapter
        ids=[h[0] for h in HEADS]
        for i,hid in enumerate(ids):
            if toc.get(hid) is None: toc[hid]=toc.get(ids[i+1],'')
    cover=fitz.open(); w,h=432,691.2; pg=cover.new_page(width=w,height=h); pg.insert_image(fitz.Rect(0,0,w,h),filename=HERE+'/cover3.jpg')
    m=fitz.open(); m.insert_pdf(cover); m.insert_pdf(fitz.open(bp))
    for pi in range(len(m)):
        if pi<5: continue
        p=m[pi]; r=p.rect; p.insert_text((r.width/2-6,r.height-38),str(pi+1),fontsize=9,fontname='tiro',color=(.35,.4,.47))
    m.set_metadata({'title':f'{TITLE}: {SUB}','author':AUTHOR,'subject':'A human account of kidney disease, waiting, and care','keywords':'kidney disease, dialysis, transplant, patient-centered care'})
    m.save(f'{OUT}/What-We-Promised.pdf',deflate=True,garbage=3)
    blank=[i+1 for i in range(1,len(m)) if len(m[i].get_text().strip())<25]
    print('PDF pages',len(m),'blank-ish pages:',blank)

def build_epub():
    css=CSS+'\nbody{text-align:left}\n.toc div{display:block}\n.toc .dots{display:none}\n'
    body=render(None,epub=True)
    chunks=re.split(r'(?=<div class="opener">)',body); files=[('front','Title and contents',chunks[0])]
    for ch in chunks[1:]:
        m=re.search(r'id="(s\d+)">(.*?)</h[12]>',ch); files.append((m.group(1),html.unescape(m.group(2)),ch))
    z=zipfile.ZipFile(f'{OUT}/What-We-Promised.epub','w'); z.writestr('mimetype','application/epub+zip',compress_type=zipfile.ZIP_STORED)
    z.writestr('META-INF/container.xml','<?xml version="1.0"?><container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>')
    z.writestr('OEBPS/style.css',css)
    X=lambda t,i:f'<?xml version="1.0" encoding="utf-8"?><!DOCTYPE html><html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops" lang="en"><head><meta charset="utf-8"/><title>{html.escape(t)}</title><link rel="stylesheet" type="text/css" href="style.css"/></head><body>{i}</body></html>'
    z.writestr('OEBPS/cover.xhtml',X('Cover','<div style="text-align:center"><img src="images/cover.jpg" alt="Cover of What We Promised" style="max-width:100%;height:auto"/></div>')); z.write(HERE+'/cover3.jpg','OEBPS/images/cover.jpg')
    for k in ALT: z.write(f'{FIG}/{k}.png',f'OEBPS/fig/{k}.png')
    items=['<item id="cover" href="cover.xhtml" media-type="application/xhtml+xml"/>','<item id="css" href="style.css" media-type="text/css"/>','<item id="coverimg" href="images/cover.jpg" media-type="image/jpeg" properties="cover-image"/>','<item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>']+[f'<item id="img_{k}" href="fig/{k}.png" media-type="image/png"/>' for k in ALT]
    spine=['<itemref idref="cover" linear="no"/>','<itemref idref="nav"/>']; nav=[]
    for fid,title,inner in files:
        z.writestr(f'OEBPS/{fid}.xhtml',X(title,inner)); items.append(f'<item id="{fid}" href="{fid}.xhtml" media-type="application/xhtml+xml"/>'); spine.append(f'<itemref idref="{fid}"/>')
        if fid!='front': nav.append(f'<li><a href="{fid}.xhtml">{html.escape(smart(title))}</a></li>')
    z.writestr('OEBPS/nav.xhtml',X('Contents',f'<nav epub:type="toc" id="toc"><h1>Contents</h1><ol>{"".join(nav)}</ol></nav>'))
    uid='urn:uuid:'+str(uuid.UUID('7a1e7de0-2026-4b1a-9d00-70ad11e5f0a2'))
    opf=f'<?xml version="1.0" encoding="utf-8"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="uid"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:identifier id="uid">{uid}</dc:identifier><dc:title>{TITLE}: {html.escape(SUB)}</dc:title><dc:creator>{AUTHOR}</dc:creator><dc:language>en-US</dc:language><dc:publisher>Cold Ischemia Foundation</dc:publisher><dc:date>2026-10-09</dc:date><dc:subject>Kidney disease; Dialysis; Transplantation; Patient-centered care</dc:subject><meta property="dcterms:modified">2026-10-09T00:00:00Z</meta></metadata><manifest>{"".join(items)}</manifest><spine>{"".join(spine)}</spine></package>'
    z.writestr('OEBPS/content.opf',opf); z.close(); print('EPUB written',len(files),'files')

def build_docx():
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    d=Document(); s=d.sections[0]; s.page_width=Inches(6); s.page_height=Inches(9); s.left_margin=s.right_margin=Inches(0.9); s.top_margin=Inches(0.9); s.bottom_margin=Inches(0.95)
    st=d.styles; N=st['Normal']; N.font.name='Georgia'; N.font.size=Pt(11); N.element.rPr.rFonts.set(qn('w:eastAsia'),'Georgia')
    N.paragraph_format.space_after=Pt(8); N.paragraph_format.line_spacing=1.35; N.paragraph_format.first_line_indent=Inches(0.25)
    def mk(name,size=11,bold=False,italic=False,color=None,align=None,before=0,after=8,indent0=True):
        x=st.add_style(name,WD_STYLE_TYPE.PARAGRAPH); x.base_style=st['Normal']; x.font.size=Pt(size); x.font.bold=bold; x.font.italic=italic
        if color: x.font.color.rgb=RGBColor.from_string(color)
        if align is not None: x.paragraph_format.alignment=align
        x.paragraph_format.space_before=Pt(before); x.paragraph_format.space_after=Pt(after)
        if indent0: x.paragraph_format.first_line_indent=Inches(0)
        return x
    for nm,sz,aft,bef,col in (('Heading 1',22,14,0,'1C2B3A'),('Heading 2',24,16,0,'1C2B3A'),('Heading 3',13,6,16,'1C2B3A')):
        x=st[nm]; x.font.name='Georgia'; x.font.size=Pt(sz); x.font.bold=True; x.font.color.rgb=RGBColor.from_string(col); x.element.rPr.rFonts.set(qn('w:ascii'),'Georgia'); x.element.rPr.rFonts.set(qn('w:hAnsi'),'Georgia')
        x.paragraph_format.space_before=Pt(bef); x.paragraph_format.space_after=Pt(aft); x.paragraph_format.first_line_indent=Inches(0); x.paragraph_format.keep_with_next=True
    mk('Part Label',12,True,color='C96A14',align=WD_ALIGN_PARAGRAPH.CENTER,after=36); mk('Chapter Label',10.5,True,color='0E8FB0',after=4)
    mk('Epigraph',10.5,italic=True,color='3A4756',align=WD_ALIGN_PARAGRAPH.LEFT,after=2); st['Epigraph'].paragraph_format.left_indent=Inches(0.4); st['Epigraph'].paragraph_format.right_indent=Inches(0.3)
    mk('EpigraphBy',9,color='6B7684',align=WD_ALIGN_PARAGRAPH.RIGHT,after=22)
    mk('Figure Caption',9,italic=True,color='4A5563',align=WD_ALIGN_PARAGRAPH.LEFT,after=14); mk('Body First',align=WD_ALIGN_PARAGRAPH.LEFT)
    mk('Reference',9.5,align=WD_ALIGN_PARAGRAPH.LEFT,after=6); st['Reference'].paragraph_format.left_indent=Inches(0.3); st['Reference'].paragraph_format.first_line_indent=Inches(-0.3)
    mk('Centered',align=WD_ALIGN_PARAGRAPH.CENTER); mk('Note Text',10,align=WD_ALIGN_PARAGRAPH.LEFT,after=4)
    mk('TOC Part',10.5,True,color='C96A14',align=WD_ALIGN_PARAGRAPH.LEFT,after=2,before=8); mk('TOC Chapter',10.5,align=WD_ALIGN_PARAGRAPH.LEFT,after=2); st['TOC Chapter'].paragraph_format.left_indent=Inches(0.25)
    mk('TOC Front',10.5,True,align=WD_ALIGN_PARAGRAPH.LEFT,after=2)
    bid=[100]
    def bookmark(p,name):
        bid[0]+=1; a=OxmlElement('w:bookmarkStart'); a.set(qn('w:id'),str(bid[0])); a.set(qn('w:name'),name); e=OxmlElement('w:bookmarkEnd'); e.set(qn('w:id'),str(bid[0])); p._p.insert(1,a); p._p.append(e)
    def link(p,text,anchor):
        h=OxmlElement('w:hyperlink'); h.set(qn('w:anchor'),anchor); r=OxmlElement('w:r'); t=OxmlElement('w:t'); t.text=text; t.set(qn('xml:space'),'preserve'); r.append(t); h.append(r); p._p.append(h)
    def addruns(p,text):
        for t,f in runs(text):
            r=p.add_run(t); r.bold=True if f=='b' else r.bold; r.italic=True if f=='i' else r.italic
    def shade(p):
        pr=p._p.get_or_add_pPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto'); sh.set(qn('w:fill'),'F6EFE3'); pr.append(sh)
        bd=OxmlElement('w:pBdr'); l=OxmlElement('w:left'); l.set(qn('w:val'),'single'); l.set(qn('w:sz'),'24'); l.set(qn('w:space'),'6'); l.set(qn('w:color'),'C96A14'); bd.append(l); pr.append(bd); p.paragraph_format.left_indent=Inches(0.15)
    pc=d.add_paragraph(style='Centered'); pc.paragraph_format.space_after=Pt(0); pc.add_run().add_picture(HERE+'/cover3.jpg',width=Inches(4.3))
    d.core_properties.title=f'{TITLE}: {SUB}'; d.core_properties.author=AUTHOR
    first=False; pend=None; hn=0; refs=False; pbnext=True
    def pbp(style,text=None,break_=True,**kw):
        p=d.add_paragraph(style=style);
        if break_: p.paragraph_format.page_break_before=True
        if text: p.add_run(text)
        return p
    for b in BL:
        t=b[0]
        if t=='part': pend=(b[1],b[2]); continue
        if t=='special':
            tag,body=b[1],b[2]
            if tag=='TITLEPAGE':
                p=pbp('Centered'); p.paragraph_format.space_before=Pt(130)
                r=p.add_run(TITLE.upper()); r.bold=True; r.font.size=Pt(28); r.font.color.rgb=RGBColor(0x1c,0x2b,0x42)
                p=d.add_paragraph(style='Centered'); p.paragraph_format.space_before=Pt(18); r=p.add_run(SUB); r.italic=True; r.font.size=Pt(13)
                p=d.add_paragraph(style='Centered'); p.paragraph_format.space_before=Pt(50); r=p.add_run(AUTHOR.upper()); r.bold=True; r.font.size=Pt(15)
                p=d.add_paragraph(style='Centered'); r=p.add_run('Cold Ischemia Foundation'); r.font.size=Pt(10)
            elif tag=='COPYRIGHT':
                for i,c in enumerate(COPY):
                    p=d.add_paragraph(style='Body First')
                    if i==0: p.paragraph_format.page_break_before=True; p.paragraph_format.space_before=Pt(150)
                    for part in re.split(r'(<strong>.*?</strong>)',c):
                        if part.startswith('<strong>'): r=p.add_run(part[8:-9]); r.bold=True
                        elif part: p.add_run(part)
                    for r in p.runs: r.font.size=Pt(9)
            elif tag=='DEDICATION':
                p=pbp('Centered'); p.paragraph_format.space_before=Pt(110); r=p.add_run(smart(' '.join(body))); r.italic=True; r.font.size=Pt(12)
            elif tag=='EPIGRAPH':
                p=d.add_paragraph(style='Centered'); p.paragraph_format.space_before=Pt(90); r=p.add_run(smart(body[0])); r.italic=True; r.font.size=Pt(10.5)
                p=d.add_paragraph(style='Centered'); r=p.add_run(smart(body[1]) if len(body)>1 else ''); r.font.size=Pt(9)
            elif tag=='TOC':
                p=d.add_paragraph(style='Heading 1'); p.paragraph_format.page_break_before=True; p.add_run('Contents'); bookmark(p,'toc')
                for hid,lvl,txt,tgt in HEADS:
                    st_='TOC Part' if txt.startswith('Part ') else ('TOC Chapter' if lvl==2 else 'TOC Front')
                    p=d.add_paragraph(style=st_); link(p,smart(txt),tgt)
            continue
        if t=='h1':
            hn+=1; k=kind(b[1]); refs=b[1].startswith('Appendix B')
            if pend:
                p=d.add_paragraph(style='Part Label'); p.paragraph_format.page_break_before=True; p.paragraph_format.space_before=Pt(96); p.add_run(f'PART {pend[0].upper()}');
                p.add_run('\n'+smart(pend[1]).upper()); bookmark(p,f'p{hn}'); pend=None; extra=False
            else: extra=True
            h=d.add_paragraph(style='Heading 2' if k=='chapter' else 'Heading 1')
            if extra: h.paragraph_format.page_break_before=True; h.paragraph_format.space_before=Pt(96)
            h.add_run(smart(b[1])); bookmark(h,f's{hn}'); first=True; continue
        if t=='h3': p=d.add_paragraph(style='Heading 3'); p.add_run(smart(b[1])); first=True; continue
        if t=='epi':
            p=d.add_paragraph(style='Epigraph'); p.add_run(smart(b[1])); p=d.add_paragraph(style='EpigraphBy'); p.add_run(smart(b[2])); continue
        if t=='fig':
            p=d.add_paragraph(style='Centered'); p.paragraph_format.keep_with_next=True; pic=p.add_run().add_picture(f'{FIG}/{b[1]}.png',width=Inches(4.6))
            pic._inline.docPr.set('descr',ALT[b[1]]); pic._inline.docPr.set('title',ALT[b[1]][:60])
            c=d.add_paragraph(style='Figure Caption'); addruns(c,b[2]); continue
        if t=='note':
            for x in b[2]: p=d.add_paragraph(style='Note Text'); addruns(p,x); shade(p)
            continue
        if t=='table':
            rows=b[1]; tb=d.add_table(rows=len(rows),cols=len(rows[0])); tb.style='Table Grid'
            for i,row in enumerate(rows):
                for j,c in enumerate(row):
                    cell=tb.cell(i,j); cell.text=''; p=cell.paragraphs[0]; p.paragraph_format.first_line_indent=Inches(0); p.paragraph_format.space_after=Pt(2); p.alignment=WD_ALIGN_PARAGRAPH.LEFT; addruns(p,c)
                    for r in p.runs:
                        r.font.size=Pt(8.5)
                        if i==0: r.bold=True; r.font.color.rgb=RGBColor(255,255,255)
                    if i==0:
                        pr=cell._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto'); sh.set(qn('w:fill'),'1C2B3A'); pr.append(sh)
            d.add_paragraph().paragraph_format.space_after=Pt(6); continue
        if t=='p':
            p=d.add_paragraph(style='Reference' if refs else ('Body First' if first else 'Normal')); first=False; addruns(p,b[1])
    ORD=['pStyle','keepNext','keepLines','pageBreakBefore','framePr','widowControl','numPr','suppressLineNumbers','pBdr','shd','tabs','suppressAutoHyphens','kinsoku','wordWrap','overflowPunct','topLinePunct','autoSpaceDE','autoSpaceDN','bidi','adjustRightInd','snapToGrid','spacing','ind','contextualSpacing','mirrorIndents','suppressOverlap','jc','textDirection','textAlignment','textboxTightWrap','outlineLvl','divId','cnfStyle','rPr','sectPr','pPrChange']
    loc=lambda e:e.tag.split('}')[1]
    for ppr in list(d.element.body.iter(qn('w:pPr')))+list(d.styles.element.iter(qn('w:pPr'))):
        ks=list(ppr); ss=sorted(ks,key=lambda e:ORD.index(loc(e)) if loc(e) in ORD else 99)
        if ks!=ss:
            for k in ks: ppr.remove(k)
            for k in ss: ppr.append(k)
    d.save(f'{OUT}/What-We-Promised.docx'); print('DOCX written')

def words():
    W=lambda s:len(re.findall(r"[A-Za-z0-9’'\-]+",s)); tot=0; main=0; inb=False
    for b in BL:
        n=0
        if b[0] in('p','h1','h3'): n=W(b[1])
        elif b[0]=='epi': n=W(b[1]+' '+b[2])
        elif b[0]=='note': n=sum(W(x) for x in b[2])
        elif b[0]=='table': n=sum(W(c) for r in b[1] for c in r)
        if b[0]=='h1':
            if b[1].startswith('Prologue'): inb=True
            if b[1].startswith('Appendix A'): inb=False
        tot+=n; main+=n if inb else 0
    return tot,main
if __name__=='__main__':
    which=sys.argv[2:] or ['docx','epub','pdf']; print('WORDS total, main text:',words())
    if 'docx' in which: build_docx()
    if 'epub' in which: build_epub()
    if 'pdf' in which: build_pdf()
